import openpyxl, pandas as pd, numpy as np, json, re
from pathlib import Path
from docx import Document
from scipy import stats
try: import pycountry
except: pycountry=None
p='② 附件二_数据表_Jolie FemTech.xlsx';wbv=openpyxl.load_workbook(p,data_only=True);wbf=openpyxl.load_workbook(p,data_only=False)
ws=wbv['跨国底稿']; h=[ws.cell(4,c).value for c in range(1,20)];rows=[]
for r in range(5,172):
 if ws.cell(r,1).value:rows.append(dict(zip(h,[ws.cell(r,c).value for c in range(1,20)])))
df=pd.DataFrame(rows)
# formula-based sheets expected vs cached check
checks=[]
def add(sheet,cell,actual,expected,tol=1e-9):
 ok=(actual==expected) if isinstance(expected,str) else (pd.notna(actual) and abs(float(actual)-float(expected))<=tol*max(1,abs(float(expected))))
 checks.append((sheet,cell,actual,expected,ok))
# descriptions
cols={'MMR':'C','log(MMR)（公式）':'N','原始GDP':'E','logGDP（公式）':'G','政治体制代码':'H','CPR':'K','CPR年份':'L'}
w=wbv['描述统计']
for i,(name,col) in enumerate(cols.items(),5):
 x=pd.to_numeric(df[name],errors='coerce').dropna().to_numpy();q=np.quantile(x,[.25,.75])
 vals=[len(x),np.mean(x),np.std(x,ddof=1),np.median(x),q[1]-q[0],np.min(x),np.max(x)]
 for c,ex in zip(range(2,9),vals):add('描述统计',w.cell(i,c).coordinate,w.cell(i,c).value,ex)
# missing
w=wbv['缺失分析']
for r,code in zip(range(5,9),[0,1,2,3]):
 x=df[df['政治体制代码']==code]; vals=[len(x),x['CPR'].notna().sum(),x['CPR'].isna().sum(),x['CPR'].isna().mean()]
 for c,ex in zip(range(3,7),vals):add('缺失分析',w.cell(r,c).coordinate,w.cell(r,c).value,ex)
# corr
w=wbv['相关系数']; pairs=[('政治体制代码','log(MMR)（公式）'),('政治体制代码','MMR'),('政治体制代码','CPR'),('logGDP（公式）','log(MMR)（公式）'),('logGDP（公式）','MMR'),('logGDP（公式）','CPR'),('政治体制代码','logGDP（公式）')]
for r,(a,b) in zip(range(5,12),pairs):
 z=df[[a,b]].dropna();add('相关系数',w.cell(r,3).coordinate,w.cell(r,3).value,len(z));add('相关系数',w.cell(r,4).coordinate,w.cell(r,4).value,stats.pearsonr(z[a],z[b]).statistic)
# independent model match including terms in full table input models JSON
models=json.loads(Path('audit_artifacts/independent_models.json').read_text())
mapping={'M1':'M1 Primary: log(MMR), ordinal regime','M2':'M2 Sensitivity: raw MMR, ordinal regime','M3':'M3 Primary: CPR 2010-2023 latest','M4':'M4 CPR + observation year','M5':'M5 CPR restricted to 2015-2023','M6':'M6 log(MMR), regime categorical','M7':'M7 CPR, regime categorical'}
term_map={'政治体制代码_n':'regime_code','logGDP（公式）_n':'log_gdp','cyearc':'cpr_year_centered_2018','d1':'regime_1','d2':'regime_2','d3':'regime_3'}
w=wbv['模型输出与HC3结果']; out=[]
for row in w.iter_rows(min_row=5,values_only=True):
 if isinstance(row[0],str) and row[0].startswith('M'):out.append(row)
for m in models:
 for rr in m['rows']:
  term=term_map.get(rr['term'],rr['term'])
  found=[x for x in out if x[0]==mapping[m['label']] and x[3]==term]
  if not found:
   checks.append(('模型输出与HC3结果','MISSING',None,(mapping[m['label']],term),False));continue
  x=found[0]; expected=[m['n'],m['k'],m['dof'],rr['beta'],rr['ols_se'],rr['ols_t'],rr['ols_p'],rr['ols_ci_l'],rr['ols_ci_u'],rr['hc3_se'],rr['hc3_t'],rr['hc3_p'],rr['hc3_ci_l'],rr['hc3_ci_u'],m['r2'],m['adjr2']]
  for pos,ex in zip([4,5,6,7,8,9,10,11,12,13,14,15,16,17,18,19],expected):add('模型输出与HC3结果',f'{mapping[m["label"]]}:{term}:col{pos+1}',x[pos],ex,tol=1e-8)
print('checks',len(checks),'fail',sum(not x[-1] for x in checks))
for x in checks:
 if not x[-1]:print('FAIL',x)
# document checks basic
D=Document('① 论文终稿_Jolie FemTech.docx');text='\n'.join(p.text for p in D.paragraphs)
print('doc_paras',len(D.paragraphs),'tables',len(D.tables),'shapes',len(D.inline_shapes))
for token in ['[Lastname]','[Institution Name]','[student email]']:print('placeholder',token,token in text)
# ISO validation
if pycountry:
 bad=[]
 for iso in df['ISO3']:
  if not pycountry.countries.get(alpha_3=iso):bad.append(iso)
 print('pycountry_bad_iso',bad)
else: print('pycountry unavailable')
