from pathlib import Path
import openpyxl, numpy as np, pandas as pd, json, re, os, math, hashlib, zipfile, glob, subprocess, sys
from scipy import stats
from statsmodels.stats.outliers_influence import variance_inflation_factor
from collections import Counter

XLSX=Path('② 附件二_数据表_Jolie FemTech.xlsx')
OUT=Path('audit_artifacts'); OUT.mkdir(exist_ok=True)
wbv=openpyxl.load_workbook(XLSX,data_only=True,read_only=False)
wbf=openpyxl.load_workbook(XLSX,data_only=False,read_only=False,keep_links=True)
ws=wbv['跨国底稿']; wsf=wbf['跨国底稿']
headers=[ws.cell(4,c).value for c in range(1,20)]
print('HEADERS',headers)
records=[]
for r in range(5,172):
    vals=[ws.cell(r,c).value for c in range(1,20)]
    if vals[0] is None: continue
    rec=dict(zip(headers,vals));rec['_row']=r
    records.append(rec)
df=pd.DataFrame(records)
df.to_csv(OUT/'cross_country_raw_from_xlsx.csv',index=False)
print('\n-- CORE DATA --')
print(df.to_string(index=False))
print('\nrows',len(df),'unique iso',df[headers[0]].nunique(),'duplicates',df[df[headers[0]].duplicated(keep=False)][[headers[0],headers[1], '_row']].to_dict('records'))
print('\n-- FORMULA AND CACHE CHECK --')
formula_issues=[]
for wsname in wbf.sheetnames:
    wf=wbf[wsname];wv=wbv[wsname]
    for row in wf.iter_rows():
        for cf in row:
            if isinstance(cf.value,str) and cf.value.startswith('='):
                cv=wv[cf.coordinate].value
                if cv is None:
                    formula_issues.append((wsname,cf.coordinate,cf.value,'MISSING_CACHED_VALUE'))
print('formula_count',sum(1 for sh in wbf for row in sh.iter_rows() for c in row if isinstance(c.value,str) and c.value.startswith('=')))
print('formula cache missing',formula_issues[:100], 'count',len(formula_issues))
# formula and cell error strings, comments, hidden sheet/rows/columns
errors=[]; comments=[]; hidden=[]
for sh in wbf.worksheets:
    if sh.sheet_state!='visible': hidden.append((sh.title,'sheet',sh.sheet_state))
    for idx,dim in sh.row_dimensions.items():
        if dim.hidden: hidden.append((sh.title,'row',idx))
    for idx,dim in sh.column_dimensions.items():
        if dim.hidden: hidden.append((sh.title,'col',idx))
    for row in sh.iter_rows():
        for c in row:
            if c.data_type=='e' or (isinstance(c.value,str) and c.value in {'#REF!','#DIV/0!','#VALUE!','#N/A','#NAME?','#NUM!'}):errors.append((sh.title,c.coordinate,c.value))
            if c.comment: comments.append((sh.title,c.coordinate,c.comment.text[:200]))
print('errors',errors,'comments',comments,'hidden',hidden)
# map columns by friendly letters expected
iso='ISO3';country='国家'; mmr='MMR'; gdp='原始GDP'; code='政治体制代码'; cpr='CPR'; cyear='CPR年份'; logg='logGDP（公式）'; logm='log(MMR)（公式）'; flag='CPR缺失标记（公式）'
# print headers if mismatch
need=[iso,country,mmr,gdp,code,cpr,cyear,logg,logm,flag]
print('missing expected', [x for x in need if x not in df.columns])
# derived validation
num=lambda c:pd.to_numeric(df[c],errors='coerce')
for col in [mmr,gdp,code,cpr,cyear,logg,logm]:df[col+'_n']=num(col)
# derived differences
print('\n-- DERIVED CHECKS --')
print('logGDP mismatch',df.loc[np.abs(df[logg+'_n']-np.log(df[gdp+'_n']))>1e-10,[iso,country,gdp,logg,'_row']].to_string(index=False))
print('logMMR mismatch',df.loc[np.abs(df[logm+'_n']-np.log(df[mmr+'_n']))>1e-10,[iso,country,mmr,logm,'_row']].to_string(index=False))
print('code invalid',df.loc[df[code+'_n'].notna() & ~df[code+'_n'].isin([0,1,2,3]),[iso,country,code,'_row']].to_string(index=False))
# Reconstruct labels and dummies positions based headers
print('all headers mapping')
for x in headers: print(repr(x))
# assess country/values reasonable
print('range mmr',df[mmr+'_n'].min(),df[mmr+'_n'].max(),'gdp',df[gdp+'_n'].min(),df[gdp+'_n'].max(),'cpr',df[cpr+'_n'].min(),df[cpr+'_n'].max(),'cpryear',df[cyear+'_n'].min(),df[cyear+'_n'].max())
# destruct statistics function
vars0=[mmr,logm,gdp,logg,code,cpr,cyear]
print('\n-- DESCRIPTIVES INDEPENDENT --')
for x in vars0:
 v=df[x+'_n'].dropna().to_numpy(); q=np.quantile(v,[.25,.5,.75])
 print(x,dict(n=len(v),mean=float(np.mean(v)),sd=float(np.std(v,ddof=1)),q1=float(q[0]),median=float(q[1]),q3=float(q[2]),iqr=float(q[2]-q[0]),min=float(np.min(v)),max=float(np.max(v))))
# pairwise correlations
pairs=[(code,logm),(code,mmr),(code,cpr),(logg,logm),(logg,mmr),(logg,cpr),(code,logg)]
print('\n-- CORR INDEPENDENT --')
for a,b in pairs:
 z=df[[a+'_n',b+'_n']].dropna(); print(a,b,'n',len(z),'r',stats.pearsonr(z.iloc[:,0],z.iloc[:,1]).statistic)
# missing
print('\n-- MISSING INDEPENDENT --')
print(df[flag].value_counts(dropna=False))
print('cpr missing n',df[cpr+'_n'].isna().sum(),'raw observed',df[cpr+'_n'].notna().sum(),'cpr model',df[[cpr+'_n',code+'_n']].dropna().shape[0])
for v in [0,1,2,3]:
 x=df[df[code+'_n']==v];miss=x[cpr+'_n'].isna().sum();print(v,'total',len(x),'missing',miss,'observed',len(x)-miss,'rate',miss/len(x))
print('missing rows')
print(df.loc[df[cpr+'_n'].isna(),[iso,country,code,flag,'_row']].to_string(index=False))
y=df.loc[df[cpr+'_n'].notna(),cyear+'_n'].to_numpy(); print('CPR year',dict(min=float(y.min()),q1=float(np.quantile(y,.25)),median=float(np.median(y)),q3=float(np.quantile(y,.75)),iqr=float(np.quantile(y,.75)-np.quantile(y,.25)),max=float(y.max())))
# OLS independent

def fit(label, ycol, predcols, subset):
    d=subset[[iso,ycol]+predcols].dropna().copy()
    y=d[ycol].to_numpy(float); X=np.c_[np.ones(len(d)),d[predcols].to_numpy(float)]; names=['Intercept']+predcols
    n,k=X.shape; dof=n-k
    xtxi=np.linalg.inv(X.T@X); beta=xtxi@X.T@y; resid=y-X@beta
    mse=(resid@resid)/dof; se=np.sqrt(np.diag(mse*xtxi)); h=np.sum((X@xtxi)*X,axis=1)
    vcovhc3=xtxi@((X.T*(resid**2/(1-h)**2))@X)@xtxi; sehc3=np.sqrt(np.diag(vcovhc3))
    tcrit=stats.t.ppf(.975,dof); t0=beta/se;t3=beta/sehc3;p0=2*stats.t.sf(np.abs(t0),dof);p3=2*stats.t.sf(np.abs(t3),dof)
    sst=((y-y.mean())**2).sum();r2=1-(resid@resid)/sst; adj=1-(1-r2)*(n-1)/dof
    print('\nMODEL',label,'n',n,'k',k,'dof',dof,'r2',r2,'adj',adj)
    for a,b,c,d1,e,f,g,h1 in zip(names,beta,se,t0,p0,sehc3,t3,p3):print(a,'beta',b,'SE',c,'t',d1,'p',e,'HC3SE',f,'HC3t',g,'HC3p',h1,'CI',b-tcrit*c,b+tcrit*c,'HC3CI',b-tcrit*f,b+tcrit*f)
    return {'label':label,'n':n,'k':k,'dof':dof,'r2':r2,'adjr2':adj,'rows':[(names[i],beta[i],se[i],t0[i],p0[i],sehc3[i],t3[i],p3[i],beta[i]-tcrit*se[i],beta[i]+tcrit*se[i],beta[i]-tcrit*sehc3[i],beta[i]+tcrit*sehc3[i])for i in range(k)],'iso':d[iso].tolist()}
base=df.copy()
base['d1']=(base[code+'_n']==1).astype(int);base['d2']=(base[code+'_n']==2).astype(int);base['d3']=(base[code+'_n']==3).astype(int);base['cyearc']=base[cyear+'_n']-2018
models=[]
models.append(fit('M1',logm+'_n',[code+'_n',logg+'_n'],base))
models.append(fit('M2',mmr+'_n',[code+'_n',logg+'_n'],base))
models.append(fit('M3',cpr+'_n',[code+'_n',logg+'_n'],base))
models.append(fit('M4',cpr+'_n',[code+'_n',logg+'_n','cyearc'],base))
models.append(fit('M5',cpr+'_n',[code+'_n',logg+'_n'],base[base[cyear+'_n']>=2015]))
models.append(fit('M6',logm+'_n',['d1','d2','d3',logg+'_n'],base[base[code+'_n'].notna()]))
models.append(fit('M7',cpr+'_n',['d1','d2','d3',logg+'_n'],base[base[code+'_n'].notna()]))
# VIF model covariates on M1 sample (regime and logGDP)
z=base[[code+'_n',logg+'_n']].dropna().to_numpy(float);print('\nVIF M1', [variance_inflation_factor(z,i) for i in range(z.shape[1])], 'r',np.corrcoef(z,rowvar=False)[0,1], 'equiv',1/(1-np.corrcoef(z,rowvar=False)[0,1]**2))
# Output full JSON models
serial=[]
for m in models:
 q=m.copy();q['rows']=[dict(zip(['term','beta','ols_se','ols_t','ols_p','hc3_se','hc3_t','hc3_p','ols_ci_l','ols_ci_u','hc3_ci_l','hc3_ci_u'],r))for r in q['rows']];serial.append(q)
(OUT/'independent_models.json').write_text(json.dumps(serial,ensure_ascii=False,indent=2))
# workbook output compare: exact expected names map from model output table
out=wbv['模型输出与HC3结果']
print('\n-- MODEL OUTPUT WORKBOOK --')
for row in out.iter_rows(min_row=4,values_only=True):
 if any(v is not None for v in row): print(row)
# xlsx package artifacts
with zipfile.ZipFile(XLSX) as z:
 names=z.namelist();
 print('\nZIP external links',[n for n in names if 'externalLink' in n]); print('calcChain', [n for n in names if 'calcChain' in n]); print('comments', [n for n in names if 'comment' in n.lower()]); print('customXML',[n for n in names if 'customXml' in n])
